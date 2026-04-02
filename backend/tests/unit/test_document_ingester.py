"""Unit tests for the Document Ingester component."""

from __future__ import annotations

import struct
import zlib
from unittest.mock import patch, MagicMock

import fitz
import pytest

from app.document_ingester import ingest, SUPPORTED_TYPES
from app.errors import FileReadError, UnsupportedFileError
from app.models import DocumentInput, DocumentModel


# --- Helpers ---


def _make_minimal_png(width: int = 2, height: int = 2, color: tuple = (255, 0, 0)) -> bytes:
    """Create a minimal valid PNG image in memory."""
    raw_data = b""
    for _ in range(height):
        raw_data += b"\x00"  # filter byte
        for _ in range(width):
            raw_data += bytes(color)

    def _chunk(chunk_type: bytes, data: bytes) -> bytes:
        c = chunk_type + data
        crc = struct.pack(">I", zlib.crc32(c) & 0xFFFFFFFF)
        return struct.pack(">I", len(data)) + c + crc

    ihdr_data = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + _chunk(b"IHDR", ihdr_data)
        + _chunk(b"IDAT", zlib.compress(raw_data))
        + _chunk(b"IEND", b"")
    )


# --- Fixtures ---


@pytest.fixture
def simple_pdf(tmp_path):
    """Create a simple PDF with text."""
    path = str(tmp_path / "test.pdf")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello from PDF")
    doc.save(path)
    doc.close()
    return path


@pytest.fixture
def simple_docx(tmp_path):
    """Create a simple Word document with text."""
    import docx

    path = str(tmp_path / "test.docx")
    doc = docx.Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    doc.save(path)
    return path


@pytest.fixture
def docx_with_image(tmp_path):
    """Create a Word document with an embedded image."""
    import docx
    from docx.shared import Inches

    path = str(tmp_path / "with_image.docx")
    doc = docx.Document()
    doc.add_paragraph("Document with image")

    # Write a small PNG to disk, then embed it
    img_path = tmp_path / "embed.png"
    img_path.write_bytes(_make_minimal_png())
    doc.add_picture(str(img_path), width=Inches(1))
    doc.save(path)
    return path


@pytest.fixture
def png_image(tmp_path):
    """Create a PNG image file."""
    path = tmp_path / "diagram.png"
    path.write_bytes(_make_minimal_png())
    return str(path)


@pytest.fixture
def jpg_image(tmp_path):
    """Create a fake JPG image file (just bytes for pass-through)."""
    path = tmp_path / "photo.jpg"
    path.write_bytes(b"\xff\xd8\xff\xe0" + b"\x00" * 100)
    return str(path)


@pytest.fixture
def svg_image(tmp_path):
    """Create an SVG image file."""
    path = tmp_path / "diagram.svg"
    path.write_bytes(b'<svg xmlns="http://www.w3.org/2000/svg"></svg>')
    return str(path)


# --- PDF Ingestion Tests ---


class TestPdfIngestion:
    """Tests for PDF document ingestion."""

    def test_pdf_returns_document_model(self, simple_pdf):
        doc_input = DocumentInput(file_path=simple_pdf, file_type="pdf")
        result = ingest(doc_input)
        assert isinstance(result, DocumentModel)
        assert result.file_type == "pdf"
        assert result.source_path == simple_pdf

    def test_pdf_extracts_text(self, simple_pdf):
        doc_input = DocumentInput(file_path=simple_pdf, file_type="pdf")
        result = ingest(doc_input)
        assert len(result.text_content) >= 1
        assert any("Hello from PDF" in t for t in result.text_content)

    def test_pdf_not_found_raises_file_read_error(self):
        doc_input = DocumentInput(file_path="/nonexistent/file.pdf", file_type="pdf")
        with pytest.raises(FileReadError):
            ingest(doc_input)


# --- DOCX Ingestion Tests ---


class TestDocxIngestion:
    """Tests for Word document ingestion."""

    def test_docx_returns_document_model(self, simple_docx):
        doc_input = DocumentInput(file_path=simple_docx, file_type="docx")
        result = ingest(doc_input)
        assert isinstance(result, DocumentModel)
        assert result.file_type == "docx"
        assert result.source_path == simple_docx

    def test_docx_extracts_paragraphs(self, simple_docx):
        doc_input = DocumentInput(file_path=simple_docx, file_type="docx")
        result = ingest(doc_input)
        assert "First paragraph" in result.text_content
        assert "Second paragraph" in result.text_content

    def test_docx_extracts_images(self, docx_with_image):
        doc_input = DocumentInput(file_path=docx_with_image, file_type="docx")
        result = ingest(doc_input)
        assert len(result.images) >= 1
        assert all(isinstance(img, bytes) for img in result.images)
        assert all(len(img) > 0 for img in result.images)

    def test_docx_not_found_raises_file_read_error(self):
        doc_input = DocumentInput(file_path="/nonexistent/file.docx", file_type="docx")
        with pytest.raises(FileReadError, match="not found"):
            ingest(doc_input)

    def test_corrupted_docx_raises_file_read_error(self, tmp_path):
        path = str(tmp_path / "bad.docx")
        with open(path, "wb") as f:
            f.write(b"not a real docx file")
        doc_input = DocumentInput(file_path=path, file_type="docx")
        with pytest.raises(FileReadError):
            ingest(doc_input)


# --- Image Ingestion Tests ---


class TestImageIngestion:
    """Tests for image file pass-through ingestion."""

    def test_png_returns_document_model(self, png_image):
        doc_input = DocumentInput(file_path=png_image, file_type="png")
        result = ingest(doc_input)
        assert isinstance(result, DocumentModel)
        assert result.file_type == "png"
        assert result.source_path == png_image

    def test_png_has_single_image(self, png_image):
        doc_input = DocumentInput(file_path=png_image, file_type="png")
        result = ingest(doc_input)
        assert len(result.images) == 1
        assert len(result.images[0]) > 0

    def test_png_has_no_text(self, png_image):
        doc_input = DocumentInput(file_path=png_image, file_type="png")
        result = ingest(doc_input)
        assert result.text_content == []

    def test_jpg_pass_through(self, jpg_image):
        doc_input = DocumentInput(file_path=jpg_image, file_type="jpg")
        result = ingest(doc_input)
        assert len(result.images) == 1
        assert result.file_type == "jpg"

    def test_svg_pass_through(self, svg_image):
        doc_input = DocumentInput(file_path=svg_image, file_type="svg")
        result = ingest(doc_input)
        assert len(result.images) == 1
        assert result.file_type == "svg"

    def test_image_not_found_raises_file_read_error(self):
        doc_input = DocumentInput(file_path="/nonexistent/image.png", file_type="png")
        with pytest.raises(FileReadError, match="not found"):
            ingest(doc_input)

    def test_empty_image_raises_file_read_error(self, tmp_path):
        path = tmp_path / "empty.png"
        path.write_bytes(b"")
        doc_input = DocumentInput(file_path=str(path), file_type="png")
        with pytest.raises(FileReadError, match="empty"):
            ingest(doc_input)

    def test_image_bytes_match_file_content(self, png_image):
        with open(png_image, "rb") as f:
            expected = f.read()
        doc_input = DocumentInput(file_path=png_image, file_type="png")
        result = ingest(doc_input)
        assert result.images[0] == expected


# --- Unsupported File Type Tests ---


class TestUnsupportedFileType:
    """Tests for unsupported file type handling."""

    def test_unsupported_type_raises_error(self):
        # DocumentInput validates file_type via Literal, so we bypass with a mock
        doc_input = MagicMock()
        doc_input.file_type = "xlsx"
        doc_input.file_path = "/some/file.xlsx"
        with pytest.raises(UnsupportedFileError):
            ingest(doc_input)

    def test_unsupported_error_lists_supported_types(self):
        doc_input = MagicMock()
        doc_input.file_type = "xlsx"
        doc_input.file_path = "/some/file.xlsx"
        with pytest.raises(UnsupportedFileError) as exc_info:
            ingest(doc_input)
        assert exc_info.value.supported_types == SUPPORTED_TYPES
