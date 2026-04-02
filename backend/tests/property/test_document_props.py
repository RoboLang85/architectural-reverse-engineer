"""Property-based tests for PDF text and image extraction round-trip.

# Feature: architectural-reverse-engineer, Property 4: PDF text and image extraction round-trip

**Validates: Requirements 2.1**

For any programmatically generated PDF containing known text blocks and embedded
images, the PDF Processor should extract text content that contains all the
original text blocks and image data that matches the original embedded images.
"""

from __future__ import annotations

import tempfile
import os

import fitz  # PyMuPDF
from hypothesis import given, settings, assume
from hypothesis import strategies as st

from app.pdf_processor import extract


# ---------------------------------------------------------------------------
# Strategies
# ---------------------------------------------------------------------------

# Generate non-empty printable ASCII text strings suitable for PDF insertion.
# We restrict to basic ASCII printable characters (letters, digits, punctuation,
# space) because PDF text rendering may alter or drop exotic Unicode whitespace
# and special characters. This keeps the round-trip property testable.
_printable_text = st.text(
    alphabet=st.characters(
        min_codepoint=32,
        max_codepoint=126,
    ),
    min_size=3,
    max_size=80,
).filter(lambda s: s.strip())

# Generate a list of 1-3 text blocks to place on PDF pages.
_text_blocks = st.lists(_printable_text, min_size=1, max_size=3)

# Generate RGB color tuples for image rectangles (0-255 per channel).
_rgb_color = st.tuples(
    st.integers(min_value=0, max_value=255),
    st.integers(min_value=0, max_value=255),
    st.integers(min_value=0, max_value=255),
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _create_pdf_with_text(text_blocks: list[str], tmp_dir: str) -> str:
    """Create a PDF with one page per text block, return the file path."""
    path = os.path.join(tmp_dir, "test.pdf")
    doc = fitz.open()
    for text in text_blocks:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    doc.save(path)
    doc.close()
    return path


def _create_png_image(color: tuple[int, int, int], tmp_dir: str, name: str) -> str:
    """Create a small solid-color PNG image, return the file path."""
    img_path = os.path.join(tmp_dir, name)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20), 0)
    pix.set_rect(pix.irect, color)
    pix.save(img_path)
    pix = None
    return img_path


def _create_pdf_with_text_and_images(
    text_blocks: list[str], colors: list[tuple[int, int, int]], tmp_dir: str
) -> str:
    """Create a PDF with text blocks and embedded colored rectangle images."""
    path = os.path.join(tmp_dir, "test_with_images.pdf")
    doc = fitz.open()

    for i, text in enumerate(text_blocks):
        page = doc.new_page()
        page.insert_text((72, 72), text)

        # Embed an image on each page if we have a color for it
        if i < len(colors):
            img_path = _create_png_image(colors[i], tmp_dir, f"img_{i}.png")
            rect = fitz.Rect(72, 120, 172, 220)
            page.insert_image(rect, filename=img_path)

    doc.save(path)
    doc.close()
    return path


# ---------------------------------------------------------------------------
# Property 4: PDF text and image extraction round-trip
# ---------------------------------------------------------------------------


@settings(max_examples=100, deadline=None)
@given(text_blocks=_text_blocks)
def test_pdf_text_extraction_contains_all_original_text(text_blocks):
    """Extracted text blocks contain all original text inserted into the PDF.

    # Feature: architectural-reverse-engineer, Property 4: PDF text and image extraction round-trip
    **Validates: Requirements 2.1**
    """
    with tempfile.TemporaryDirectory() as tmp_dir:
        pdf_path = _create_pdf_with_text(text_blocks, tmp_dir)
        result = extract(pdf_path)

        # Join all extracted text into one string for substring matching
        all_extracted = " ".join(result.text_blocks)

        for original_text in text_blocks:
            # PDF text extraction strips leading/trailing whitespace, so
            # normalise the original before comparison.
            normalised = original_text.strip()
            assert normalised in all_extracted, (
                f"Original text block not found in extracted content.\n"
                f"  Missing: {normalised!r}\n"
                f"  Extracted: {all_extracted!r}"
            )


@settings(max_examples=100, deadline=None)
@given(
    text_blocks=_text_blocks,
    colors=st.lists(_rgb_color, min_size=1, max_size=3),
)
def test_pdf_image_extraction_returns_nonempty_images(text_blocks, colors):
    """Extracted images are non-empty for PDFs with embedded images.

    # Feature: architectural-reverse-engineer, Property 4: PDF text and image extraction round-trip
    **Validates: Requirements 2.1**
    """
    with tempfile.TemporaryDirectory() as tmp_dir:
        pdf_path = _create_pdf_with_text_and_images(text_blocks, colors, tmp_dir)
        result = extract(pdf_path)

        # Number of embedded images = min(len(text_blocks), len(colors))
        expected_image_count = min(len(text_blocks), len(colors))

        assert len(result.images) >= expected_image_count, (
            f"Expected at least {expected_image_count} images, "
            f"got {len(result.images)}"
        )

        for i, img_bytes in enumerate(result.images):
            assert isinstance(img_bytes, bytes), f"Image {i} is not bytes"
            assert len(img_bytes) > 0, f"Image {i} is empty"


# ---------------------------------------------------------------------------
# Property 5: Word document text and image extraction round-trip
# ---------------------------------------------------------------------------
# Feature: architectural-reverse-engineer, Property 5: Word document text and image extraction round-trip
#
# **Validates: Requirements 2.2**
#
# For any programmatically generated Word document containing known text
# paragraphs and embedded images, the Analyzer should extract text content
# that contains all the original paragraphs and image data that matches the
# original embedded images.
# ---------------------------------------------------------------------------

import io
import struct
import zlib

import docx as python_docx
from docx.shared import Inches

from app import document_ingester
from app.models import DocumentInput


# --- Strategies for Property 5 ---

# Non-empty printable text paragraphs (ASCII letters/digits/punctuation/space).
_docx_paragraph = st.text(
    alphabet=st.characters(min_codepoint=32, max_codepoint=126),
    min_size=3,
    max_size=80,
).filter(lambda s: s.strip())

_docx_paragraphs = st.lists(_docx_paragraph, min_size=1, max_size=5)


def _make_minimal_png(width: int = 4, height: int = 4, color: tuple[int, int, int] = (255, 0, 0)) -> bytes:
    """Create a minimal valid PNG image in memory (no external deps)."""

    def _chunk(chunk_type: bytes, data: bytes) -> bytes:
        c = chunk_type + data
        return struct.pack(">I", len(data)) + c + struct.pack(">I", zlib.crc32(c) & 0xFFFFFFFF)

    header = b"\x89PNG\r\n\x1a\n"
    ihdr_data = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    ihdr = _chunk(b"IHDR", ihdr_data)

    raw_rows = b""
    for _ in range(height):
        raw_rows += b"\x00"  # filter byte
        for _ in range(width):
            raw_rows += bytes(color)

    idat = _chunk(b"IDAT", zlib.compress(raw_rows))
    iend = _chunk(b"IEND", b"")
    return header + ihdr + idat + iend


def _create_docx_with_text(paragraphs: list[str], tmp_dir: str) -> str:
    """Create a .docx file with the given paragraphs, return the file path."""
    path = os.path.join(tmp_dir, "test.docx")
    doc = python_docx.Document()
    for para in paragraphs:
        doc.add_paragraph(para)
    doc.save(path)
    return path


def _create_docx_with_text_and_images(
    paragraphs: list[str],
    image_count: int,
    tmp_dir: str,
) -> tuple[str, list[bytes]]:
    """Create a .docx with paragraphs and embedded PNG images.

    Returns (docx_path, list_of_original_png_bytes).
    """
    path = os.path.join(tmp_dir, "test_with_images.docx")
    doc = python_docx.Document()

    original_images: list[bytes] = []
    for i, para in enumerate(paragraphs):
        doc.add_paragraph(para)
        if i < image_count:
            # Vary color per image so they are distinct
            color = ((50 * (i + 1)) % 256, (80 * (i + 1)) % 256, (120 * (i + 1)) % 256)
            png_bytes = _make_minimal_png(width=4, height=4, color=color)
            original_images.append(png_bytes)
            doc.add_picture(io.BytesIO(png_bytes), width=Inches(0.5))

    doc.save(path)
    return path, original_images


@settings(max_examples=100, deadline=None)
@given(paragraphs=_docx_paragraphs)
def test_docx_text_extraction_contains_all_original_paragraphs(paragraphs):
    """Extracted text_content contains every original paragraph.

    # Feature: architectural-reverse-engineer, Property 5: Word document text and image extraction round-trip
    **Validates: Requirements 2.2**
    """
    with tempfile.TemporaryDirectory() as tmp_dir:
        docx_path = _create_docx_with_text(paragraphs, tmp_dir)
        doc_input = DocumentInput(file_path=docx_path, file_type="docx")
        result = document_ingester.ingest(doc_input)

        all_extracted = " ".join(result.text_content)

        for original in paragraphs:
            stripped = original.strip()
            assert stripped in all_extracted, (
                f"Original paragraph not found in extracted content.\n"
                f"  Missing: {stripped!r}\n"
                f"  Extracted: {all_extracted!r}"
            )


@settings(max_examples=100, deadline=None)
@given(
    paragraphs=_docx_paragraphs,
    image_count=st.integers(min_value=1, max_value=3),
)
def test_docx_image_extraction_returns_nonempty_images(paragraphs, image_count):
    """Extracted images are non-empty bytes for Word docs with embedded images.

    # Feature: architectural-reverse-engineer, Property 5: Word document text and image extraction round-trip
    **Validates: Requirements 2.2**
    """
    with tempfile.TemporaryDirectory() as tmp_dir:
        docx_path, original_images = _create_docx_with_text_and_images(
            paragraphs, image_count, tmp_dir
        )
        doc_input = DocumentInput(file_path=docx_path, file_type="docx")
        result = document_ingester.ingest(doc_input)

        expected_count = min(len(paragraphs), image_count)

        assert len(result.images) >= expected_count, (
            f"Expected at least {expected_count} images, got {len(result.images)}"
        )

        for i, img_bytes in enumerate(result.images):
            assert isinstance(img_bytes, bytes), f"Image {i} is not bytes"
            assert len(img_bytes) > 0, f"Image {i} is empty"


# ---------------------------------------------------------------------------
# Property 6: Unsupported file type error includes supported types list
# ---------------------------------------------------------------------------
# Feature: architectural-reverse-engineer, Property 6: Unsupported file type error includes supported types list
#
# **Validates: Requirements 2.4**
#
# For any file with an extension not in the supported set (pdf, docx, png,
# jpg, svg), the Analyzer should return an error that lists all supported
# file types.
# ---------------------------------------------------------------------------

from unittest.mock import MagicMock

from app.errors import UnsupportedFileError

SUPPORTED_TYPES = {"pdf", "docx", "png", "jpg", "svg"}

# Strategy: generate non-empty lowercase alpha strings that are NOT in the supported set.
_unsupported_ext = st.text(
    alphabet=st.characters(min_codepoint=ord("a"), max_codepoint=ord("z")),
    min_size=1,
    max_size=10,
).filter(lambda s: s not in SUPPORTED_TYPES)


@settings(max_examples=100, deadline=None)
@given(ext=_unsupported_ext)
def test_unsupported_file_type_error_includes_supported_types(ext):
    """Unsupported file type raises UnsupportedFileError listing all supported types.

    # Feature: architectural-reverse-engineer, Property 6: Unsupported file type error includes supported types list
    **Validates: Requirements 2.4**
    """
    # Bypass Pydantic Literal validation by using a MagicMock
    mock_input = MagicMock()
    mock_input.file_type = ext
    mock_input.file_path = f"/tmp/fake_file.{ext}"

    try:
        document_ingester.ingest(mock_input)
        # Should not reach here
        raise AssertionError(
            f"Expected UnsupportedFileError for extension '{ext}', but no error was raised."
        )
    except UnsupportedFileError as err:
        # Verify the error's supported_types list contains ALL supported types
        assert set(err.supported_types) == SUPPORTED_TYPES, (
            f"Error supported_types mismatch.\n"
            f"  Expected: {SUPPORTED_TYPES}\n"
            f"  Got: {set(err.supported_types)}"
        )
        # Verify the error message mentions the unsupported type
        assert ext in err.message, (
            f"Error message should mention the unsupported type '{ext}'.\n"
            f"  Message: {err.message}"
        )
